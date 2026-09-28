"""Business glossary (Atlas ``GlossaryService``).

Glossaries, terms and categories are stored as entities of the types
``AtlasGlossary`` / ``AtlasGlossaryTerm`` / ``AtlasGlossaryCategory`` with the
relationships of ``models/0000-Area0/0011-glossary_model.json`` - exactly like
Atlas - so they get versioning, audits, classifications (which propagate from a term
to the entities it is assigned to) and search for free.  This module converts
between those entities and the glossary REST objects.
"""
from __future__ import annotations

import copy
import csv
import zipfile
import io
import json
import logging
from typing import Any, Dict, Iterable, List, Optional, Tuple

from ..errors import AtlasBaseException, AtlasErrorCode
from ..safety import check_zip, neutralize_formula, restore_formula_escape
from ..repository.entity_store import EntityStore, MEANING_REL

log = logging.getLogger(__name__)

GLOSSARY = "AtlasGlossary"
TERM = "AtlasGlossaryTerm"
CATEGORY = "AtlasGlossaryCategory"
TERM_ANCHOR = "AtlasGlossaryTermAnchor"
CATEGORY_ANCHOR = "AtlasGlossaryCategoryAnchor"
TERM_CATEGORIZATION = "AtlasGlossaryTermCategorization"
CATEGORY_HIERARCHY = "AtlasGlossaryCategoryHierarchyLink"

INVALID_NAME_CHARS = ("@", ".", "<", ">")

# attribute name -> (relationship type, is end2 attribute)   (AtlasGlossaryTerm.Relation)
TERM_RELATIONS = {
    "seeAlso": ("AtlasGlossaryRelatedTerm", False),
    "synonyms": ("AtlasGlossarySynonym", False),
    "antonyms": ("AtlasGlossaryAntonym", False),
    "preferredToTerms": ("AtlasGlossaryPreferredTerm", True),
    "preferredTerms": ("AtlasGlossaryPreferredTerm", False),
    "replacementTerms": ("AtlasGlossaryReplacementTerm", True),
    "replacedBy": ("AtlasGlossaryReplacementTerm", False),
    "translationTerms": ("AtlasGlossaryTranslation", True),
    "translatedTerms": ("AtlasGlossaryTranslation", False),
    "isA": ("AtlasGlossaryIsARelationship", True),
    "classifies": ("AtlasGlossaryIsARelationship", False),
    "validValues": ("AtlasGlossaryValidValue", True),
    "validValuesFor": ("AtlasGlossaryValidValue", False),
}

IMPORT_HEADERS = ["GlossaryName", "TermName", "ShortDescription", "LongDescription", "Examples", "Abbreviation",
                  "Usage", "AdditionalAttributes", "TranslationTerms", "ValidValuesFor", "Synonyms", "ReplacedBy",
                  "ValidValues", "ReplacementTerms", "SeeAlso", "TranslatedTerms", "IsA", "Antonyms", "Classifies",
                  "PreferredToTerms", "PreferredTerms"]
IMPORT_RELATION_COLUMNS = ["translationTerms", "validValuesFor", "synonyms", "replacedBy", "validValues",
                           "replacementTerms", "seeAlso", "translatedTerms", "isA", "antonyms", "classifies",
                           "preferredToTerms", "preferredTerms"]
AUDIT_HEADERS = ["Record Type", "Name", "Glossary Name", "Short Description", "Long Description", "Status",
                 "Classifications", "Custom Attributes", "Related Categories / Parent", "Qualified Name", "GUID"]

GLOSSARY_SETTABLE = ("name", "shortDescription", "longDescription", "language", "usage")
TERM_SETTABLE = ("name", "shortDescription", "longDescription", "abbreviation", "usage")
CATEGORY_SETTABLE = ("name", "shortDescription", "longDescription")


def _nn(d: dict) -> dict:
    """Drop None values (Atlas serialises glossary objects with NON_NULL)."""
    return {k: v for k, v in d.items() if v is not None}


def _name_invalid(name: Optional[str]) -> bool:
    return bool(name) and any(c in name for c in INVALID_NAME_CHARS)


def _active(items) -> List[dict]:
    if not items:
        return []
    if isinstance(items, dict):
        items = [items]
    return [i for i in items if i and i.get("relationshipStatus", "ACTIVE") == "ACTIVE"
            and i.get("entityStatus", "ACTIVE") == "ACTIVE"]


def _rel_attrs(item: dict) -> dict:
    ra = item.get("relationshipAttributes") or {}
    return ra.get("attributes") or {} if "attributes" in ra else ra


def _sort_page(items: List[dict], key: str, sort: str, offset: int, limit: int) -> List[dict]:
    items = sorted(items, key=lambda x: (x.get(key) or ""), reverse=str(sort).upper() == "DESC")
    offset = max(0, int(offset or 0))
    limit = int(limit if limit is not None else -1)
    return items[offset:] if limit < 0 else items[offset:offset + limit]


class GlossaryService:
    def __init__(self, entities: EntityStore):
        self.entities = entities
        self.store = entities.store

    # ================================================================== loading
    async def _load(self, guid: str, type_name: str) -> dict:
        if not guid:
            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, "guid is null/empty")
        try:
            ent = (await self.entities.get_by_guid(guid, min_ext_info=True))["entity"]
        except AtlasBaseException:
            raise AtlasBaseException(AtlasErrorCode.INSTANCE_GUID_NOT_FOUND, guid)
        if ent["typeName"] != type_name:
            raise AtlasBaseException(AtlasErrorCode.INSTANCE_GUID_NOT_FOUND, guid)
        if ent.get("status") == "DELETED":
            raise AtlasBaseException(AtlasErrorCode.INSTANCE_GUID_DELETED, guid)
        return ent

    async def _parents_of(self, category_guids: Iterable[str]) -> Dict[str, str]:
        guids = list(dict.fromkeys(category_guids))
        if not guids:
            return {}
        q = {"bool": {"filter": [{"term": {"typeName": CATEGORY_HIERARCHY}}, {"term": {"status": "ACTIVE"}},
                                 {"terms": {"end2Guid": guids}}]}}
        out = {}
        async for _, r in self.store.scan(self.store.relationships, q):
            out[r["end2Guid"]] = r["end1Guid"]
        return out

    @staticmethod
    def _term_header(item: dict) -> dict:
        a = _rel_attrs(item)
        return _nn({"termGuid": item["guid"], "relationGuid": item.get("relationshipGuid"),
                    "displayText": item.get("displayText"), "description": a.get("description"),
                    "expression": a.get("expression"), "steward": a.get("steward"), "source": a.get("source"),
                    "status": a.get("status"), "qualifiedName": item.get("qualifiedName")})

    @staticmethod
    def _category_header(item: dict, parents: Dict[str, str]) -> dict:
        a = _rel_attrs(item)
        return _nn({"categoryGuid": item["guid"], "parentCategoryGuid": parents.get(item["guid"]),
                    "relationGuid": item.get("relationshipGuid"), "displayText": item.get("displayText"),
                    "description": a.get("description")})

    @staticmethod
    def _glossary_header(item: Optional[dict]) -> Optional[dict]:
        if not item:
            return None
        return _nn({"glossaryGuid": item["guid"], "relationGuid": item.get("relationshipGuid"),
                    "displayText": item.get("displayText")})

    @staticmethod
    def _base(ent: dict) -> dict:
        a = ent.get("attributes") or {}
        out = {"guid": ent["guid"], "qualifiedName": a.get("qualifiedName"), "name": a.get("name"),
               "shortDescription": a.get("shortDescription"), "longDescription": a.get("longDescription"),
               "additionalAttributes": a.get("additionalAttributes")}
        if ent.get("classifications"):
            out["classifications"] = ent["classifications"]
        return out

    async def _glossary_dto(self, ent: dict) -> dict:
        a = ent.get("attributes") or {}
        ra = ent.get("relationshipAttributes") or {}
        out = self._base(ent)
        out.update({"language": a.get("language"), "usage": a.get("usage")})
        terms = _active(ra.get("terms"))
        cats = _active(ra.get("categories"))
        if terms:
            out["terms"] = [self._term_header(t) for t in terms]
        if cats:
            parents = await self._parents_of(c["guid"] for c in cats)
            out["categories"] = [self._category_header(c, parents) for c in cats]
        return _nn(out)

    def _term_dto(self, ent: dict) -> dict:
        a = ent.get("attributes") or {}
        ra = ent.get("relationshipAttributes") or {}
        out = self._base(ent)
        out.update({"examples": a.get("examples"), "abbreviation": a.get("abbreviation"), "usage": a.get("usage")})
        anchor = _active(ra.get("anchor"))
        if anchor:
            out["anchor"] = self._glossary_header(anchor[0])
        assigned = _active(ra.get("assignedEntities"))
        if assigned:
            out["assignedEntities"] = assigned
        cats = _active(ra.get("categories"))
        if cats:
            out["categories"] = [_nn({"categoryGuid": c["guid"], "relationGuid": c.get("relationshipGuid"),
                                      "description": _rel_attrs(c).get("description"),
                                      "displayText": c.get("displayText"), "status": _rel_attrs(c).get("status")})
                                 for c in cats]
        for attr in TERM_RELATIONS:
            items = _active(ra.get(attr))
            if items:
                out[attr] = [self._term_header(i) for i in items]
        return _nn(out)

    async def _category_dto(self, ent: dict) -> dict:
        ra = ent.get("relationshipAttributes") or {}
        out = self._base(ent)
        anchor = _active(ra.get("anchor"))
        if anchor:
            out["anchor"] = self._glossary_header(anchor[0])
        parent = _active(ra.get("parentCategory"))
        children = _active(ra.get("childrenCategories"))
        parents = await self._parents_of([p["guid"] for p in parent] + [c["guid"] for c in children])
        if parent:
            out["parentCategory"] = self._category_header(parent[0], parents)
        if children:
            out["childrenCategories"] = [self._category_header(c, parents) for c in children]
        terms = _active(ra.get("terms"))
        if terms:
            out["terms"] = [self._term_header(t) for t in terms]
        return _nn(out)

    # ================================================================== glossaries
    async def get_glossaries(self, limit: int = -1, offset: int = 0, sort: str = "ASC") -> List[dict]:
        q = {"bool": {"filter": [{"term": {"typeName": GLOSSARY}}, {"term": {"status": "ACTIVE"}}]}}
        guids = []
        async for g, d in self.store.scan(self.store.entities, q, source=["guid", "displayText"]):
            guids.append((d.get("displayText") or "", g))
        guids.sort(reverse=str(sort).upper() == "DESC")
        offset = max(0, int(offset or 0))
        sel = guids[offset:] if int(limit) < 0 else guids[offset:offset + int(limit)]
        return [await self.get_glossary(g) for _, g in sel]

    async def get_glossary(self, guid: str) -> dict:
        return await self._glossary_dto(await self._load(guid, GLOSSARY))

    async def get_detailed_glossary(self, guid: str) -> dict:
        g = await self.get_glossary(guid)
        g["termInfo"] = {}
        g["categoryInfo"] = {}
        for t in g.get("terms", []):
            g["termInfo"][t["termGuid"]] = await self.get_term(t["termGuid"])
        for c in g.get("categories", []):
            g["categoryInfo"][c["categoryGuid"]] = await self.get_category(c["categoryGuid"])
        return g

    async def _find_guid(self, type_name: str, qualified_name: str) -> Optional[str]:
        return await self.entities.find_guid_by_unique_attributes(type_name, {"qualifiedName": qualified_name},
                                                                  include_subtypes=False)

    async def create_glossary(self, glossary: dict, user: str) -> dict:
        if not glossary:
            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, "Glossary definition missing")
        g = dict(glossary)
        if not g.get("qualifiedName"):
            if not g.get("name"):
                raise AtlasBaseException(AtlasErrorCode.GLOSSARY_QUALIFIED_NAME_CANT_BE_DERIVED)
            if _name_invalid(g["name"]):
                raise AtlasBaseException(AtlasErrorCode.INVALID_DISPLAY_NAME)
            g["qualifiedName"] = g["name"]
        if await self._find_guid(GLOSSARY, g["qualifiedName"]):
            raise AtlasBaseException(AtlasErrorCode.GLOSSARY_ALREADY_EXISTS, g.get("name") or g["qualifiedName"])
        entity = {"typeName": GLOSSARY, "guid": "-1",
                  "attributes": {k: g.get(k) for k in ("qualifiedName", "name", "shortDescription", "longDescription",
                                                       "language", "usage", "additionalAttributes")}}
        if g.get("classifications"):
            entity["classifications"] = g["classifications"]
        res = await self.entities.create_or_update({"entity": entity}, user)
        return await self.get_glossary(res["guidAssignments"]["-1"])

    async def update_glossary(self, guid: str, glossary: dict, user: str) -> dict:
        if not glossary:
            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, "Glossary is null/empty")
        if not glossary.get("name"):
            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, "DisplayName can't be null/empty")
        if _name_invalid(glossary["name"]):
            raise AtlasBaseException(AtlasErrorCode.INVALID_DISPLAY_NAME)
        cur = await self._load(guid, GLOSSARY)
        attrs = {k: glossary.get(k) for k in ("name", "shortDescription", "longDescription", "language", "usage",
                                               "additionalAttributes")}
        await self.entities.create_or_update({"entity": {"typeName": GLOSSARY, "guid": cur["guid"], "attributes": attrs}}, user)
        return await self.get_glossary(guid)

    async def partial_update_glossary(self, guid: str, updates: Dict[str, str], user: str) -> dict:
        if not updates:
            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, "PartialUpdates missing or empty")
        g = await self.get_glossary(guid)
        for k, v in updates.items():
            if k not in GLOSSARY_SETTABLE:
                raise AtlasBaseException(AtlasErrorCode.INVALID_PARTIAL_UPDATE_ATTR, k, "Glossary")
            g[k] = v
        return await self.update_glossary(guid, g, user)

    async def delete_glossary(self, guid: str, user: str) -> None:
        g = await self.get_glossary(guid)
        for t in g.get("terms", []):
            await self.delete_term(t["termGuid"], user)
        for c in g.get("categories", []):
            try:
                await self.delete_category(c["categoryGuid"], user)
            except AtlasBaseException as e:
                if e.error_code not in (AtlasErrorCode.INSTANCE_GUID_DELETED, AtlasErrorCode.INSTANCE_GUID_NOT_FOUND):
                    raise
        await self.entities.delete_by_guids([guid], user)

    # ================================================================== terms
    async def get_term(self, guid: str) -> dict:
        return self._term_dto(await self._load(guid, TERM))

    async def _glossary_qn(self, glossary_guid: Optional[str]) -> str:
        if not glossary_guid:
            raise AtlasBaseException(AtlasErrorCode.INVALID_NEW_ANCHOR_GUID)
        g = await self._load(glossary_guid, GLOSSARY)
        qn = (g.get("attributes") or {}).get("qualifiedName")
        if not qn:
            raise AtlasBaseException(AtlasErrorCode.GLOSSARY_QUALIFIED_NAME_CANT_BE_DERIVED)
        return qn

    @staticmethod
    def _related_term_value(h: dict) -> dict:
        attrs = {k: h.get(k) for k in ("description", "expression", "source", "steward")}
        attrs["status"] = h.get("status") or "ACTIVE"
        return {"guid": h.get("termGuid"), "typeName": TERM, "relationshipAttributes": {"attributes": attrs}}

    def _term_relationship_attrs(self, term: dict, guid: Optional[str]) -> Dict[str, Any]:
        out: Dict[str, Any] = {}
        for attr in TERM_RELATIONS:
            items = term.get(attr) or []
            seen = {}
            for h in items:
                if not h.get("termGuid"):
                    raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, f"termGuid missing in {attr}")
                if guid and h["termGuid"] == guid:
                    raise AtlasBaseException(AtlasErrorCode.INVALID_TERM_RELATION_TO_SELF)
                if h["termGuid"] not in seen or (not seen[h["termGuid"]].get("relationGuid") and h.get("relationGuid")):
                    seen[h["termGuid"]] = h
            out[attr] = [self._related_term_value(h) for h in seen.values()]
        cats = {}
        for c in term.get("categories") or []:
            if not c.get("categoryGuid"):
                continue
            cats.setdefault(c["categoryGuid"], c)
        out["categories"] = [{"guid": c["categoryGuid"], "typeName": CATEGORY,
                              "relationshipAttributes": {"attributes": _nn({"description": c.get("description"),
                                                                            "status": c.get("status")})}}
                             for c in cats.values()]
        return out

    @staticmethod
    def _duplicate_related(term: dict) -> Optional[str]:
        for attr in TERM_RELATIONS:
            seen = set()
            for h in term.get(attr) or []:
                if h.get("termGuid") in seen:
                    return h["termGuid"]
                seen.add(h.get("termGuid"))
        return None

    async def create_term(self, term: dict, user: str) -> dict:
        if not term:
            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, "GlossaryTerm definition missing")
        if not term.get("anchor"):
            raise AtlasBaseException(AtlasErrorCode.MISSING_MANDATORY_ANCHOR)
        name = term.get("name")
        if not name:
            raise AtlasBaseException(AtlasErrorCode.GLOSSARY_TERM_QUALIFIED_NAME_CANT_BE_DERIVED)
        if _name_invalid(name):
            raise AtlasBaseException(AtlasErrorCode.INVALID_DISPLAY_NAME)
        glossary_guid = term["anchor"].get("glossaryGuid")
        qn = f"{name}@{await self._glossary_qn(glossary_guid)}"
        if await self._find_guid(TERM, qn):
            raise AtlasBaseException(AtlasErrorCode.GLOSSARY_TERM_ALREADY_EXISTS, qn)
        rel = self._term_relationship_attrs(term, None)
        rel["anchor"] = {"guid": glossary_guid, "typeName": GLOSSARY}
        entity = {"typeName": TERM, "guid": "-1",
                  "attributes": {"qualifiedName": qn, "name": name, "shortDescription": term.get("shortDescription"),
                                 "longDescription": term.get("longDescription"), "examples": term.get("examples"),
                                 "abbreviation": term.get("abbreviation"), "usage": term.get("usage"),
                                 "additionalAttributes": term.get("additionalAttributes")},
                  "relationshipAttributes": rel}
        if term.get("classifications"):
            entity["classifications"] = term["classifications"]
        res = await self.entities.create_or_update({"entity": entity}, user)
        return await self.get_term(res["guidAssignments"]["-1"])

    async def create_terms(self, terms: List[dict], user: str) -> List[dict]:
        if terms is None:
            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, "glossaryTerm(s) is null/empty")
        return [await self.create_term(t, user) for t in terms]

    async def update_term(self, guid: str, term: dict, user: str, check_duplicates: bool = True) -> dict:
        if not term:
            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, "atlasGlossaryTerm is null/empty")
        name = term.get("name")
        if not name:
            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, "DisplayName can't be null/empty")
        if _name_invalid(name):
            raise AtlasBaseException(AtlasErrorCode.INVALID_DISPLAY_NAME)
        if check_duplicates:
            dup = self._duplicate_related(term)
            if dup:
                d = await self.get_term(dup)
                raise AtlasBaseException(AtlasErrorCode.GLOSSARY_TERM_ALREADY_EXISTS, d.get("qualifiedName"))
        cur = self._term_dto(await self._load(guid, TERM))
        anchor = term.get("anchor")
        if not anchor:
            raise AtlasBaseException(AtlasErrorCode.MISSING_MANDATORY_ANCHOR)
        qn = cur.get("qualifiedName")
        cur_glossary = (cur.get("anchor") or {}).get("glossaryGuid")
        if anchor.get("glossaryGuid") != cur_glossary:
            qn = f"{cur.get('name')}@{await self._glossary_qn(anchor.get('glossaryGuid'))}"
            other = await self._find_guid(TERM, qn)
            if other and other != guid:
                raise AtlasBaseException(AtlasErrorCode.GLOSSARY_TERM_ALREADY_EXISTS, qn)
        rel = self._term_relationship_attrs(term, guid)
        rel["anchor"] = {"guid": anchor.get("glossaryGuid"), "typeName": GLOSSARY}
        attrs = {"qualifiedName": qn, "name": name}
        for k in ("shortDescription", "longDescription", "examples", "abbreviation", "usage", "additionalAttributes"):
            attrs[k] = term.get(k)
        await self.entities.create_or_update({"entity": {"typeName": TERM, "guid": guid, "attributes": attrs,
                                                         "relationshipAttributes": rel}}, user)
        if name != cur.get("name") or qn != cur.get("qualifiedName"):
            await self.entities.refresh_meanings([a["guid"] for a in cur.get("assignedEntities", [])])
        return await self.get_term(guid)

    async def partial_update_term(self, guid: str, updates: Dict[str, str], user: str) -> dict:
        if not updates:
            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, "PartialUpdates missing or empty")
        t = await self.get_term(guid)
        for k, v in updates.items():
            if k not in TERM_SETTABLE:
                raise AtlasBaseException(AtlasErrorCode.INVALID_PARTIAL_UPDATE_ATTR, "Glossary Term", k)
            t[k] = v
        return await self.update_term(guid, t, user)

    async def delete_term(self, guid: str, user: str) -> None:
        t = await self.get_term(guid)
        if t.get("assignedEntities"):
            raise AtlasBaseException(AtlasErrorCode.TERM_HAS_ENTITY_ASSOCIATION, guid, len(t["assignedEntities"]))
        await self.entities.delete_by_guids([guid], user)

    # ---- assignments
    async def get_assigned_entities(self, guid: str, limit: int, offset: int, sort: str) -> List[dict]:
        t = await self.get_term(guid)
        return _sort_page(t.get("assignedEntities", []), "displayText", sort, offset, limit)

    async def assign_term(self, guid: str, objects: List[dict], user: str) -> None:
        t = await self.get_term(guid)
        assigned = {a["guid"] for a in t.get("assignedEntities", [])}
        for o in objects or []:
            target = o.get("guid")
            if not target and o.get("typeName") and o.get("uniqueAttributes"):
                target = await self.entities.get_guid_by_unique_attributes(o["typeName"], o["uniqueAttributes"])
            if not target:
                raise AtlasBaseException(AtlasErrorCode.INVALID_OBJECT_ID, o)
            if target in assigned:
                continue
            attrs = _rel_attrs(o) if o.get("relationshipAttributes") else {}
            await self.entities.create_relationship({"typeName": MEANING_REL, "end1": {"guid": guid},
                                                     "end2": {"guid": target}, "attributes": attrs}, user)
            assigned.add(target)

    async def remove_term_assignment(self, guid: str, objects: List[dict], user: str) -> None:
        t = await self.get_term(guid)
        assigned = {a["guid"]: a for a in t.get("assignedEntities", [])}
        for o in objects or []:
            rg = o.get("relationshipGuid")
            if not rg:
                raise AtlasBaseException(AtlasErrorCode.TERM_DISSOCIATION_MISSING_RELATION_GUID)
            cur = assigned.get(o.get("guid"))
            if cur is None or cur.get("relationshipGuid") != rg:
                raise AtlasBaseException(AtlasErrorCode.INVALID_TERM_DISSOCIATION, rg, guid, o.get("guid"))
            await self.entities.delete_relationship(rg, user)

    # ---- lists
    async def get_glossary_term_headers(self, guid: str, limit: int, offset: int, sort: str) -> List[dict]:
        g = await self.get_glossary(guid)
        return _sort_page(g.get("terms", []), "displayText", sort, offset, limit)

    async def get_glossary_terms(self, guid: str, limit: int, offset: int, sort: str) -> List[dict]:
        return [await self.get_term(h["termGuid"]) for h in await self.get_glossary_term_headers(guid, limit, offset, sort)]

    async def get_glossary_category_headers(self, guid: str, limit: int, offset: int, sort: str) -> List[dict]:
        g = await self.get_glossary(guid)
        return _sort_page(g.get("categories", []), "displayText", sort, offset, limit)

    async def get_glossary_categories(self, guid: str, limit: int, offset: int, sort: str) -> List[dict]:
        return [await self.get_category(h["categoryGuid"])
                for h in await self.get_glossary_category_headers(guid, limit, offset, sort)]

    async def get_category_terms(self, guid: str, limit: int, offset: int, sort: str) -> List[dict]:
        c = await self.get_category(guid)
        return _sort_page(c.get("terms", []), "displayText", sort, offset, limit)

    async def get_related_terms(self, guid: str, limit: int, offset: int, sort: str) -> Dict[str, List[dict]]:
        t = await self.get_term(guid)
        return {attr: _sort_page(t[attr], "displayText", sort, offset, limit) for attr in TERM_RELATIONS if t.get(attr)}

    async def get_related_categories(self, guid: str, limit: int, offset: int, sort: str) -> Dict[str, List[dict]]:
        c = await self.get_category(guid)
        out = {}
        if c.get("parentCategory"):
            out["parent"] = [c["parentCategory"]]
        if c.get("childrenCategories"):
            out["children"] = _sort_page(c["childrenCategories"], "displayText", sort, offset, limit)
        return out

    # ================================================================== categories
    async def get_category(self, guid: str) -> dict:
        return await self._category_dto(await self._load(guid, CATEGORY))

    async def _category_qn(self, name: str, glossary_guid: str, parent_guid: Optional[str]) -> str:
        if parent_guid:
            parent = await self._load(parent_guid, CATEGORY)
            return f"{name}.{(parent.get('attributes') or {}).get('qualifiedName')}"
        return f"{name}@{await self._glossary_qn(glossary_guid)}"

    async def _check_children_same_glossary(self, glossary_guid: str, children: List[dict]) -> None:
        for ch in children:
            c = await self.get_category(ch.get("categoryGuid"))
            if (c.get("anchor") or {}).get("glossaryGuid") != glossary_guid:
                raise AtlasBaseException(AtlasErrorCode.INVALID_CHILD_CATEGORY_DIFFERENT_GLOSSARY, ch.get("categoryGuid"))

    async def _requalify_children(self, guid: str, user: str) -> None:
        """Recompute qualifiedName of all descendants after a parent/anchor change."""
        c = await self.get_category(guid)
        for ch in c.get("childrenCategories", []):
            child = await self._load(ch["categoryGuid"], CATEGORY)
            name = (child.get("attributes") or {}).get("name")
            qn = f"{name}.{c['qualifiedName']}"
            if (child.get("attributes") or {}).get("qualifiedName") != qn:
                await self.entities.create_or_update({"entity": {"typeName": CATEGORY, "guid": child["guid"],
                                                                 "attributes": {"qualifiedName": qn}}}, user)
            await self._requalify_children(child["guid"], user)

    def _category_relationship_attrs(self, cat: dict) -> Dict[str, Any]:
        rel: Dict[str, Any] = {}
        p = cat.get("parentCategory")
        rel["parentCategory"] = ({"guid": p["categoryGuid"], "typeName": CATEGORY,
                                  "relationshipAttributes": {"attributes": _nn({"description": p.get("description")})}}
                                 if p and p.get("categoryGuid") else None)
        rel["childrenCategories"] = [{"guid": c["categoryGuid"], "typeName": CATEGORY,
                                      "relationshipAttributes": {"attributes": _nn({"description": c.get("description")})}}
                                     for c in cat.get("childrenCategories") or [] if c.get("categoryGuid")]
        terms = []
        for t in cat.get("terms") or []:
            if not t.get("termGuid"):
                raise AtlasBaseException(AtlasErrorCode.MISSING_TERM_ID_FOR_CATEGORIZATION)
            terms.append({"guid": t["termGuid"], "typeName": TERM, "relationshipAttributes": {"attributes": _nn(
                {"description": t.get("description"), "status": t.get("status")})}})
        rel["terms"] = terms
        return rel

    async def create_category(self, cat: dict, user: str) -> dict:
        if not cat:
            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, "GlossaryCategory definition missing")
        if not cat.get("anchor"):
            raise AtlasBaseException(AtlasErrorCode.MISSING_MANDATORY_ANCHOR)
        name = cat.get("name")
        if not name:
            raise AtlasBaseException(AtlasErrorCode.GLOSSARY_CATEGORY_QUALIFIED_NAME_CANT_BE_DERIVED)
        if _name_invalid(name):
            raise AtlasBaseException(AtlasErrorCode.INVALID_DISPLAY_NAME)
        glossary_guid = cat["anchor"].get("glossaryGuid")
        parent_guid = (cat.get("parentCategory") or {}).get("categoryGuid")
        qn = await self._category_qn(name, glossary_guid, parent_guid)
        if await self._find_guid(CATEGORY, qn):
            raise AtlasBaseException(AtlasErrorCode.GLOSSARY_CATEGORY_ALREADY_EXISTS, qn)
        await self._check_children_same_glossary(glossary_guid, cat.get("childrenCategories") or [])
        rel = self._category_relationship_attrs(cat)
        rel["anchor"] = {"guid": glossary_guid, "typeName": GLOSSARY}
        entity = {"typeName": CATEGORY, "guid": "-1",
                  "attributes": {"qualifiedName": qn, "name": name, "shortDescription": cat.get("shortDescription"),
                                 "longDescription": cat.get("longDescription"),
                                 "additionalAttributes": cat.get("additionalAttributes")},
                  "relationshipAttributes": rel}
        if cat.get("classifications"):
            entity["classifications"] = cat["classifications"]
        res = await self.entities.create_or_update({"entity": entity}, user)
        guid = res["guidAssignments"]["-1"]
        await self._requalify_children(guid, user)
        return await self.get_category(guid)

    async def create_categories(self, cats: List[dict], user: str) -> List[dict]:
        if cats is None:
            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, "glossaryCategory is null/empty")
        return [await self.create_category(c, user) for c in cats]

    async def update_category(self, guid: str, cat: dict, user: str) -> dict:
        if not cat:
            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, "GlossaryCategory is null/empty")
        name = cat.get("name")
        if not name:
            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, "DisplayName can't be null/empty")
        if _name_invalid(name):
            raise AtlasBaseException(AtlasErrorCode.INVALID_DISPLAY_NAME)
        cur = await self.get_category(guid)
        anchor = cat.get("anchor") or cur.get("anchor")
        if not anchor or not anchor.get("glossaryGuid"):
            raise AtlasBaseException(AtlasErrorCode.INVALID_NEW_ANCHOR_GUID)
        glossary_guid = anchor["glossaryGuid"]
        new_parent = (cat.get("parentCategory") or {}).get("categoryGuid")
        old_parent = (cur.get("parentCategory") or {}).get("categoryGuid")
        qn = cur.get("qualifiedName")
        if new_parent != old_parent or glossary_guid != (cur.get("anchor") or {}).get("glossaryGuid"):
            qn = await self._category_qn(cur.get("name"), glossary_guid, new_parent)
            other = await self._find_guid(CATEGORY, qn)
            if other and other != guid:
                raise AtlasBaseException(AtlasErrorCode.GLOSSARY_CATEGORY_ALREADY_EXISTS, qn)
        new_children = [c for c in cat.get("childrenCategories") or []
                        if c.get("categoryGuid") not in {x["categoryGuid"] for x in cur.get("childrenCategories", [])}]
        await self._check_children_same_glossary(glossary_guid, new_children)
        rel = self._category_relationship_attrs(cat)
        rel["anchor"] = {"guid": glossary_guid, "typeName": GLOSSARY}
        attrs = {"qualifiedName": qn, "name": name}
        for k in ("shortDescription", "longDescription", "additionalAttributes"):
            attrs[k] = cat.get(k)
        removed_children = [c["categoryGuid"] for c in cur.get("childrenCategories", [])
                            if c["categoryGuid"] not in {x.get("categoryGuid") for x in cat.get("childrenCategories") or []}]
        await self.entities.create_or_update({"entity": {"typeName": CATEGORY, "guid": guid, "attributes": attrs,
                                                         "relationshipAttributes": rel}}, user)
        await self._requalify_children(guid, user)
        for ch in removed_children:
            await self._make_top_level(ch, glossary_guid, user)
        return await self.get_category(guid)

    async def _make_top_level(self, guid: str, glossary_guid: str, user: str) -> None:
        c = await self._load(guid, CATEGORY)
        name = (c.get("attributes") or {}).get("name")
        qn = f"{name}@{await self._glossary_qn(glossary_guid)}"
        await self.entities.create_or_update({"entity": {"typeName": CATEGORY, "guid": guid,
                                                         "attributes": {"qualifiedName": qn}}}, user)
        await self._requalify_children(guid, user)

    async def partial_update_category(self, guid: str, updates: Dict[str, str], user: str) -> dict:
        if not updates:
            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, "PartialUpdates missing or empty")
        c = await self.get_category(guid)
        for k, v in updates.items():
            if k not in CATEGORY_SETTABLE:
                raise AtlasBaseException(AtlasErrorCode.INVALID_PARTIAL_UPDATE_ATTR, "Glossary Category", k)
            c[k] = v
        return await self.update_category(guid, c, user)

    async def delete_category(self, guid: str, user: str) -> None:
        c = await self.get_category(guid)
        glossary_guid = (c.get("anchor") or {}).get("glossaryGuid")
        await self.entities.delete_by_guids([guid], user)
        for ch in c.get("childrenCategories", []):
            try:
                await self._make_top_level(ch["categoryGuid"], glossary_guid, user)
            except AtlasBaseException:
                pass

    # ================================================================== search / export
    async def _rows(self, glossary_guids: Optional[List[str]]) -> Tuple[List[dict], Dict[str, dict]]:
        if not glossary_guids:
            glossary_guids = [g["guid"] for g in await self.get_glossaries()]
        rows, ext = [], {}
        for gg in glossary_guids:
            try:
                g = await self.get_detailed_glossary(gg)
            except AtlasBaseException:
                continue
            ext[gg] = g
            gname = g.get("name")
            for t in g.get("termInfo", {}).values():
                rows.append({"recordType": "TERM", "guid": t["guid"], "glossaryGuid": gg, "glossaryName": gname,
                             "name": t.get("name"), "shortDescription": t.get("shortDescription"),
                             "longDescription": t.get("longDescription"), "qualifiedName": t.get("qualifiedName"),
                             "status": _term_status(t), "classifications": _fmt_classifications(t),
                             "customAttributes": _fmt_attrs(t.get("additionalAttributes")),
                             "relatedCategoriesOrParent": ", ".join(sorted(c.get("displayText") or "" for c in t.get("categories", []))),
                             "entityStatus": "ACTIVE", "examples": "|".join(t.get("examples") or []),
                             "abbreviation": t.get("abbreviation"), "usage": t.get("usage"),
                             **{attr: "|".join(f"{h.get('displayText') or h['termGuid']}:{h.get('status') or 'ACTIVE'}"
                                               for h in t.get(attr, [])) for attr in TERM_RELATIONS},
                             "_obj": t})
            for c in g.get("categoryInfo", {}).values():
                rows.append({"recordType": "CATEGORY", "guid": c["guid"], "glossaryGuid": gg, "glossaryName": gname,
                             "name": c.get("name"), "shortDescription": c.get("shortDescription"),
                             "longDescription": c.get("longDescription"), "qualifiedName": c.get("qualifiedName"),
                             "status": str((c.get("additionalAttributes") or {}).get("status") or "ACTIVE"),
                             "classifications": _fmt_classifications(c),
                             "customAttributes": _fmt_attrs(c.get("additionalAttributes")),
                             "relatedCategoriesOrParent": (c.get("parentCategory") or {}).get("displayText") or "",
                             "entityStatus": "ACTIVE", "_obj": c})
        return rows, ext

    @staticmethod
    def _filter_rows(rows: List[dict], p: dict) -> List[dict]:
        def contains(v, s):
            return not s or (s.lower() in str(v or "").lower())
        rt = str(p.get("recordType") or "ALL").upper()
        scoped = bool(p.get("glossaryGuid") or p.get("glossaryGuids"))
        out = []
        for r in rows:
            if rt != "ALL" and r["recordType"] != rt:
                continue
            if p.get("glossaryGuid") and r["glossaryGuid"] != p["glossaryGuid"]:
                continue
            if p.get("glossaryGuids") and r["glossaryGuid"] not in p["glossaryGuids"]:
                continue
            if not contains(r["status"], p.get("statusContains")):
                continue
            if not contains(r["classifications"], p.get("classificationContains")):
                continue
            q = p.get("searchQuery")
            if q and not (any(contains(r.get(k), q) for k in ("name", "shortDescription", "longDescription", "status",
                                                                 "classifications", "customAttributes",
                                                                 "relatedCategoriesOrParent"))
                          or (not scoped and contains(r.get("glossaryName"), q))):
                continue
            if p.get("excludeDeleted", True) and str(r.get("entityStatus")).upper() == "DELETED":
                continue
            out.append(r)
        return out

    @staticmethod
    def _sort_rows(rows: List[dict], sort_by: Optional[str], sort_order: Optional[str]) -> None:
        key = (sort_by or "name").lower()
        field = {"glossaryname": "glossaryName", "glossary": "glossaryName", "status": "status",
                 "qualifiedname": "qualifiedName", "recordtype": "recordType", "type": "recordType",
                 "glossarytype": "recordType"}.get(key, "name")
        rows.sort(key=lambda r: (str(r.get(field) or "").lower(), str(r.get("name") or "").lower()),
                  reverse=str(sort_order or "").upper().startswith("DESC"))

    async def search(self, p: dict) -> dict:
        p = p or {}
        glossary = p.get("glossary") or {}
        gtype = str(p.get("glossaryType") or "ALL").upper()
        fp = {"recordType": gtype, "statusContains": p.get("status"), "classificationContains": p.get("classificationContains"),
              "searchQuery": p.get("searchQuery"), "excludeDeleted": p.get("excludeDeleted", True),
              "glossaryGuid": glossary.get("guid")}
        rows, ext = await self._rows([glossary["guid"]] if glossary.get("guid") else None)
        rows = self._filter_rows(rows, fp)
        if glossary.get("name"):
            rows = [r for r in rows if glossary["name"].lower() in str(r.get("glossaryName") or "").lower()]
        self._sort_rows(rows, p.get("sortBy", "name"), p.get("sortOrder", "ASCENDING"))
        total = len(rows)
        offset = max(0, int(p.get("offset") or 0))
        limit = int(p.get("limit", 25) if p.get("limit") is not None else 25)
        page = rows[offset:] if limit <= 0 else rows[offset:offset + limit]
        details: Dict[str, dict] = {}
        for r in page:
            g = ext[r["glossaryGuid"]]
            d = details.setdefault(r["glossaryGuid"], {
                "name": g.get("name"), "guid": g["guid"], "shortDescription": g.get("shortDescription"),
                "longDescription": g.get("longDescription"),
                "status": str((g.get("additionalAttributes") or {}).get("status") or "ACTIVE"),
                "terms": [], "categories": []})
            o = r["_obj"]
            item = {"name": o.get("name"), "qualifiedName": o.get("qualifiedName"), "guid": o["guid"],
                    "shortDescription": o.get("shortDescription"), "longDescription": o.get("longDescription"),
                    "status": r["status"], "classifications": list(o.get("classifications") or []),
                    "customAttributes": o.get("additionalAttributes")}
            if r["recordType"] == "TERM" and gtype != "CATEGORY":
                d["terms"].append(item)
            elif r["recordType"] == "CATEGORY" and gtype != "TERM":
                d["categories"].append(item)
        return {"glossary": list(details.values()), "approximateCount": total}

    @staticmethod
    def export_parameters(pm: Dict[str, Any]) -> dict:
        """Atlas GlossaryExportParameterMapper.fromCreateFileRequest."""
        if not pm:
            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, "Export format (CSV or XLSX) is required")
        if pm.get("exportParameters"):
            p = dict(pm["exportParameters"])
        elif any(k in pm for k in ("glossaryType", "glossary", "limit", "offset", "status")) or \
                not any(k in pm for k in ("format", "mode", "recordType")):
            gt = str(pm.get("glossaryType") or "ALL").upper()
            p = {"format": str(pm.get("format") or "CSV").upper(), "mode": "AUDIT",
                 "excludeDeleted": bool(pm.get("excludeDeleted", False)),
                 "searchQuery": pm.get("searchQuery"), "classificationContains": pm.get("classificationContains"),
                 "statusContains": pm.get("status"), "sortBy": pm.get("sortBy", "name"),
                 "sortOrder": pm.get("sortOrder", "ASCENDING"), "recordType": gt if gt in ("TERM", "CATEGORY") else "ALL",
                 "glossaryGuid": (pm.get("glossary") or {}).get("guid")}
        else:
            p = dict(pm)
        p["format"] = str(p.get("format") or "CSV").upper()
        if p["format"] not in ("CSV", "XLSX"):
            p["format"] = "CSV"
        p["mode"] = str(p.get("mode") or "AUDIT").upper()
        return p

    async def export_rows(self, p: dict) -> Tuple[List[str], List[List[str]]]:
        guids = [p["glossaryGuid"]] if p.get("glossaryGuid") else (p.get("glossaryGuids") or None)
        rows, _ = await self._rows(guids)
        rows = self._filter_rows(rows, p)
        if str(p.get("sortBy") or "").lower() == "glossaryname":
            self._sort_rows(rows, "glossaryName", p.get("sortOrder"))
        else:
            self._sort_rows(rows, "name", p.get("sortOrder"))
        if p.get("mode") == "IMPORT_COMPATIBLE":
            data = [[r.get("glossaryName") or "", r.get("name") or "", r.get("shortDescription") or "",
                     r.get("longDescription") or "", r.get("examples") or "", r.get("abbreviation") or "",
                     r.get("usage") or "", r.get("customAttributes") or ""] + [r.get(c) or "" for c in IMPORT_RELATION_COLUMNS]
                    for r in rows if r["recordType"] == "TERM"]
            return IMPORT_HEADERS, data
        data = [[r["recordType"], r.get("name") or "", r.get("glossaryName") or "", r.get("shortDescription") or "",
                 r.get("longDescription") or "", r.get("status") or "", r.get("classifications") or "",
                 r.get("customAttributes") or "", r.get("relatedCategoriesOrParent") or "",
                 r.get("qualifiedName") or "", r["guid"]] for r in rows]
        return AUDIT_HEADERS, data

    # ================================================================== import
    async def import_terms(self, rows: List[List[str]], user: str) -> dict:
        success, failed = [], []
        glossary_guids: Dict[str, str] = {}
        created: Dict[str, str] = {}   # term qualifiedName -> guid

        def cell(r, i):
            return r[i].strip() if len(r) > i and r[i] is not None and str(r[i]).strip() != "" else None
        for n, r in enumerate(rows, start=1):
            gname, tname = cell(r, 0), cell(r, 1)
            if not gname:
                failed.append({"parentObjectName": "", "childObjectName": tname, "importStatus": "FAILED",
                               "remarks": f"The GlossaryName is blank for the record : {r}", "rowNumber": n})
                continue
            try:
                if gname not in glossary_guids:
                    gguid = await self._find_guid(GLOSSARY, gname)
                    if gguid is None:
                        gguid = (await self.create_glossary({"name": gname, "qualifiedName": gname}, user))["guid"]
                    glossary_guids[gname] = gguid
                if not tname:
                    raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST, f"The TermName is blank for provided record: {r}")
                addl = None
                if cell(r, 7):
                    addl = {}
                    for kv in cell(r, 7).split("|"):
                        parts = kv.split(":")
                        if len(parts) % 2 != 0:
                            raise AtlasBaseException(AtlasErrorCode.BAD_REQUEST,
                                                     "AdditionalAttributes needs to be a key:value pair")
                        addl[parts[0]] = parts[1]
                term = {"name": tname, "anchor": {"glossaryGuid": glossary_guids[gname]},
                        "shortDescription": cell(r, 2), "longDescription": cell(r, 3),
                        "examples": cell(r, 4).split("|") if cell(r, 4) else None, "abbreviation": cell(r, 5),
                        "usage": cell(r, 6), "additionalAttributes": addl}
                t = await self.create_term(term, user)
                created[t["qualifiedName"]] = t["guid"]
                success.append({"parentObjectName": gname, "childObjectName": tname, "importStatus": "SUCCESS",
                                "remarks": json.dumps({"termGuid": t["guid"], "qualifiedName": t["qualifiedName"]})})
            except AtlasBaseException as e:
                failed.append({"parentObjectName": gname, "childObjectName": tname, "importStatus": "FAILED",
                               "remarks": e.message, "rowNumber": n})
        if not success:
            raise AtlasBaseException(AtlasErrorCode.GLOSSARY_IMPORT_FAILED)
        # second pass: relations ("GlossaryName:TermName|...")
        for n, r in enumerate(rows, start=1):
            gname, tname = cell(r, 0), cell(r, 1)
            if not gname or not tname:
                continue
            qn = f"{tname}@{gname}"
            guid = created.get(qn) or await self._find_guid(TERM, qn)
            if guid is None:
                continue
            additions: Dict[str, List[dict]] = {}
            errors = []
            for i, attr in enumerate(IMPORT_RELATION_COLUMNS, start=8):
                v = cell(r, i)
                if not v:
                    continue
                for ref in v.split("|"):
                    parts = ref.split(":")
                    if len(parts) != 2:
                        errors.append(f"Incorrect relation data specified for the term : {tname}@{gname}")
                        continue
                    rqn = f"{parts[1]}@{parts[0]}"
                    if rqn.lower() == qn.lower():
                        errors.append("Invalid relationship specified for Term. Term cannot have a relationship with self")
                        continue
                    rg = created.get(rqn) or await self._find_guid(TERM, rqn)
                    if rg is None:
                        errors.append(f"The provided Reference {parts[1]}@{parts[0]} does not exist at Atlas referred at "
                                      f"record with TermName  : {tname} and GlossaryName : {gname}")
                        continue
                    additions.setdefault(attr, []).append({"termGuid": rg})
            if additions:
                try:
                    t = await self.get_term(guid)
                    for attr, hs in additions.items():
                        existing = {h["termGuid"] for h in t.get(attr, [])}
                        t[attr] = t.get(attr, []) + [h for h in hs if h["termGuid"] not in existing]
                    await self.update_term(guid, t, user, check_duplicates=False)
                except AtlasBaseException as e:
                    errors.append(e.message)
            if errors:
                failed.append({"parentObjectName": gname, "childObjectName": tname, "importStatus": "FAILED",
                               "remarks": "\n".join(errors), "rowNumber": n})
        return {"failedImportInfoList": failed, "successImportInfoList": success}


def _term_status(t: dict) -> str:
    for c in t.get("categories") or []:
        if c.get("status"):
            return c["status"]
    s = (t.get("additionalAttributes") or {}).get("status")
    return str(s) if s is not None else "ACTIVE"


def _fmt_classifications(o: dict) -> str:
    return ", ".join(sorted(c.get("typeName") for c in o.get("classifications") or [] if c.get("typeName")))


def _fmt_attrs(a: Optional[dict]) -> str:
    return "; ".join(f"{k}={v}" for k, v in (a or {}).items() if v is not None)


def read_tabular_file(file_name: str, data: bytes) -> List[List[str]]:
    """Atlas FileUtils.readFileData: CSV or Excel, header row skipped."""
    ext = (file_name or "").rsplit(".", 1)[-1].lower() if "." in (file_name or "") else ""
    rows: List[List[str]] = []
    if ext == "csv":
        text = data.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        header = next(reader, None)
        if not header:
            raise AtlasBaseException(AtlasErrorCode.NO_DATA_FOUND)
        rows = [[restore_formula_escape(c) for c in r] for r in reader if len(r) > 1]
    elif ext in ("xlsx", "xls"):
        if ext == "xls":
            raise AtlasBaseException(AtlasErrorCode.INVALID_FILE_TYPE, file_name)
        try:
            import openpyxl
        except ImportError:  # pragma: no cover
            raise AtlasBaseException(AtlasErrorCode.NOT_SUPPORTED, "XLSX files (install openpyxl)")
        try:
            check_zip(zipfile.ZipFile(io.BytesIO(data)), 256 * 1024 * 1024, "XLSX")
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except AtlasBaseException:
            raise
        except Exception:
            raise AtlasBaseException(AtlasErrorCode.NOT_VALID_FILE, "XLSX")
        it = wb.worksheets[0].iter_rows(values_only=True)
        next(it, None)
        for r in it:
            vals = ["" if v is None else restore_formula_escape(str(v).strip()) for v in r]
            if any(vals):
                rows.append(vals)
    else:
        raise AtlasBaseException(AtlasErrorCode.INVALID_FILE_TYPE, file_name)
    if not rows:
        raise AtlasBaseException(AtlasErrorCode.NO_DATA_FOUND)
    return rows


def write_tabular_file(path, fmt: str, headers: List[str], data: List[List[str]]) -> None:
    if fmt == "XLSX":
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Glossary Export"
        ws.append(headers)
        for row in data:
            # strings starting with "=" would become formulas in openpyxl / Excel (formula injection)
            ws.append([neutralize_formula((v or "")[:32766]) for v in row])
        wb.save(str(path))
        return
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, quoting=csv.QUOTE_ALL)
        w.writerow(headers)
        w.writerows([[neutralize_formula(v) for v in row] for row in data])
